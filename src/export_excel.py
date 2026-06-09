"""Exportación de los resultados a un libro de Excel (.xlsx).

A diferencia del CSV, el .xlsx guarda los marcadores como **texto explícito**,
así que Excel no los reinterpreta como fechas (el bug clásico de "1-0" → fecha).
Genera un único libro con tres hojas: Partidos, Grupos y Ratings, con formato
listo para usar.
"""
from __future__ import annotations

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

# Columnas que DEBEN tratarse como texto para que Excel no las convierta a fecha.
TEXT_COLUMNS = {"marcador_exacto", "top3_marcadores"}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_sheet(ws, df: pd.DataFrame, title: str):
    ws.title = title

    # Forzar a texto las columnas de marcador (evita la conversión a fecha).
    df = df.copy()
    for col in df.columns:
        if col in TEXT_COLUMNS:
            df[col] = df[col].astype(str)

    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            col_name = df.columns[c_idx - 1]
            if r_idx == 1:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_name in TEXT_COLUMNS:
                cell.number_format = "@"  # formato de texto explícito

    # Anchos de columna aproximados al contenido.
    for c_idx, col_name in enumerate(df.columns, 1):
        max_len = max(
            [len(str(col_name))]
            + [len(str(v)) for v in df[col_name].head(200).tolist()]
        )
        ws.column_dimensions[get_column_letter(c_idx)].width = min(max_len + 2, 48)

    ws.freeze_panes = "A2"  # fija la fila de encabezados
    ws.auto_filter.ref = ws.dimensions


def write_workbook(preds: pd.DataFrame, standings: pd.DataFrame,
                   ratings: pd.DataFrame, path: str):
    from openpyxl import Workbook

    wb = Workbook()
    _write_sheet(wb.active, preds, "Partidos")
    _write_sheet(wb.create_sheet(), standings, "Grupos")
    _write_sheet(wb.create_sheet(), ratings, "Ratings")
    wb.save(path)
    return path
